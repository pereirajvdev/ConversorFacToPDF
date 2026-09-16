""
""

from pathlib import Path
from .fac_decoder import decode_fac

def convert_to_docx(fac_path: Path, output_path: Path) -> None:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_LINE_SPACING

    text = decode_fac(fac_path.read_bytes())

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # Monospace keeps the spaces/indentation from the original FAC layout.
    style = doc.styles["Normal"]
    style.font.name = "Courier New"
    style.font.size = Pt(9.5)

    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)

    doc.save(output_path)